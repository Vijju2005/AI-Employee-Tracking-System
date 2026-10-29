"""
EBT System — PowerPoint Generator for Project Presentation
Generates a 12-slide presentation deck using python-pptx.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(22, 58, 95)       # #163A5F
    BLUE = RGBColor(37, 99, 235)      # #2563EB
    DARK = RGBColor(30, 41, 59)       # #1E293B
    LIGHT_BG = RGBColor(248, 250, 252)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(100, 116, 139)

    slides_data = [
        {
            "title": "Employee Break Tracking & Availability Prediction System",
            "subtitle": "EBT System | Final Project Presentation & Handover\nTechnologies: Python, Flask, MySQL, TiDB Cloud, Scikit-Learn\nAuthor: Vijju2005",
            "bullets": []
        },
        {
            "title": "1. Problem Statement",
            "subtitle": "Limitations in Modern Workforce Management",
            "bullets": [
                "Traditional punch-card and biometric systems only log start and end of shifts.",
                "Real employee availability and micro-break patterns remain completely untracked.",
                "Managers face difficulty predicting when team members will return from breaks.",
                "Manual paper/spreadsheet tracking leads to time discrepancy and dispute over overtime/working hours.",
                "Lack of predictive intelligence to optimize task assignment during active working windows."
            ]
        },
        {
            "title": "2. Objectives",
            "subtitle": "Core System Capabilities & Goals",
            "bullets": [
                "Automate server-side timestamping (UTC storage, localized IST Asia/Kolkata display).",
                "Implement dynamic working-hours calculation deducting multi-category breaks automatically.",
                "Provide multi-break session support (Tea, Lunch, Coffee, Personal, Other) with conflict prevention.",
                "Deploy machine learning availability predictions using Random Forest Regressors.",
                "Ensure strict Role-Based Access Control (Admin vs. Employee portals).",
                "Support instant enterprise reporting with ReportLab PDF and OpenPyXL Excel downloads."
            ]
        },
        {
            "title": "3. Technology Stack",
            "subtitle": "Robust, Scalable, Production-Grade Technologies",
            "bullets": [
                "Backend: Python 3.14, Flask 3.1, Gunicorn WSGI server.",
                "Database: TiDB Cloud Starter (Managed MySQL-compatible cluster) with TLS/SSL encryption.",
                "Frontend: HTML5, CSS3 (Enterprise Navy-Blue Design System), Vanilla JS, Chart.js.",
                "Machine Learning: Scikit-Learn (RandomForestRegressor, MAE/RMSE metrics), Pandas, NumPy.",
                "Reporting: ReportLab (Custom-styled PDF) and OpenPyXL (Excel .xlsx with auto-sizing).",
                "Hosting & DevOps: Render Cloud Web Service, GitHub CI/CD integration, Pytest test suite."
            ]
        },
        {
            "title": "4. System Architecture",
            "subtitle": "Clean Layered Web Architecture",
            "bullets": [
                "Presentation Layer: Responsive Jinja2 templates, Chart.js analytics, real-time JavaScript timers.",
                "Routing & Controller Layer: Flask Blueprints (auth, dashboard, employees, attendance, breaks, leave, reports, ml).",
                "Business Logic Layer: Dynamic hour calculations, break conflict detection, ML feature engineering.",
                "Data Access Layer: Flask-SQLAlchemy with connection pooling, pool pre-ping, and TLS verification.",
                "External Services: Flask-Mail notification engine, TiDB Cloud database gateway."
            ]
        },
        {
            "title": "5. Admin & Employee Modules",
            "subtitle": "Distinct Role-Based Experiences",
            "bullets": [
                "Admin Portal: 4 KPI overview cards, 7-day attendance trends, department distribution, employee directory management.",
                "Admin Leave Review: One-click Approve/Reject with manager notes and automated email status notifications.",
                "Employee Portal: Status hero card, check-in/out, start/end breaks, and dynamic working hours timer.",
                "Employee Leave Management: Apply for Casual/Sick/Annual leave with Full, Half, or Partial day types.",
                "Security: @admin_required route protection and custom HTTP error handlers (403, 404, 500)."
            ]
        },
        {
            "title": "6. Attendance & Break Tracking",
            "subtitle": "High Precision Time Accounting",
            "bullets": [
                "Check-In & Check-Out: Prevents duplicate sessions; records UTC timestamps safely.",
                "Live Working Hours: Automatically calculates elapsed net working time every second on the client dashboard.",
                "Multi-Break Lifecycle: Start and end multiple breaks with live active break status badges.",
                "Automatic Break Closing: If an employee checks out during an active break, the system closes the break cleanly.",
                "Formula: Net Working Time = (Gross Elapsed Session Duration) - (Total Completed Breaks Duration)."
            ]
        },
        {
            "title": "7. Machine Learning Availability Prediction",
            "subtitle": "Random Forest Predictive Intelligence",
            "bullets": [
                "Tier 1 (Cold Start < 5 records): Global category average duration fallback.",
                "Tier 2 (5 to 19 records): Personalized historical rolling average.",
                "Tier 3 (>= 20 records): Scikit-Learn RandomForestRegressor (50 trees, random state 42).",
                "Engineered Features: Day of week, hour of day, encoded category, scheduled daily shift duration.",
                "Accuracy Evaluation: Outputs real-time MAE (Mean Absolute Error) and RMSE (Root Mean Squared Error).",
                "Availability Formula: Estimated Available = max(0, Scheduled Shift Minutes - Predicted Break Duration)."
            ]
        },
        {
            "title": "8. Analytics & Enterprise Reports",
            "subtitle": "Actionable Insights & Audit Compliance",
            "bullets": [
                "Interactive Dashboard Charts: 7-day stacked attendance volume and department employee counts.",
                "ReportLab PDF Generation: Branded A4 landscape tables, custom header banners, alternating row fills.",
                "OpenPyXL Excel Generation: Structured workbooks, auto-fitted column widths, styled header cells.",
                "Supported Report Types: Attendance history, Break records, Leave summary, and Working hours audit."
            ]
        },
        {
            "title": "9. Testing & Quality Assurance",
            "subtitle": "100% Automated Test Suite Verification",
            "bullets": [
                "Framework: Pytest and Pytest-Flask with isolated in-memory test database.",
                "Total Tests: 37 test cases covering 10 functional modules.",
                "Pass Rate: 100% passing (37 passed in 5.07 seconds).",
                "Coverage: Blueprint registration, model methods, CSRF verification, RBAC route protection, open redirect defenses, ML pipeline, and report generation."
            ]
        },
        {
            "title": "10. Cloud Deployment & DevOps",
            "subtitle": "Live Production Infrastructure",
            "bullets": [
                "Production URL: https://ai-employee-tracking-system.onrender.com",
                "Cloud Database: TiDB Cloud Starter (gateway01.ap-southeast-1.prod.aws.tidbcloud.com:4000).",
                "Database Health: HTTP 200 OK, 8 initialized tables, TLS/SSL connection active.",
                "Web Server: Gunicorn with multi-worker support and 120-second request timeout.",
                "GitHub Repository: Vijju2005/AI-Employee-Tracking-System (main branch in sync)."
            ]
        },
        {
            "title": "11. Conclusion & Future Scope",
            "subtitle": "Summary and Roadmap",
            "bullets": [
                "Project Summary: Delivered a complete, secure, responsive, and ML-powered workforce tracking system.",
                "Production Ready: Tested, verified, and deployed on live HTTPS cloud infrastructure.",
                "Future Scope 1: Mobile application with GPS / Geofencing check-in validation.",
                "Future Scope 2: Facial recognition biometric verification on web punch-in.",
                "Future Scope 3: Integration with enterprise calendar APIs (Google Calendar / Microsoft Teams)."
            ]
        }
    ]

    blank_layout = prs.slide_layouts[6]

    for idx, data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)

        # Background header band
        top_bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.3))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = NAVY
        top_bar.line.color.rgb = NAVY

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(24 if idx > 0 else 26)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # Subtitle
        if data.get("subtitle"):
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.8))
            stf = sub_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = data["subtitle"]
            sp.font.size = Pt(16)
            sp.font.color.rgb = BLUE
            sp.font.bold = True

        # Bullets / Content
        if data["bullets"]:
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.8))
            btf = body_box.text_frame
            btf.word_wrap = True
            for b_idx, bullet in enumerate(data["bullets"]):
                bp = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
                bp.text = f"•  {bullet}"
                bp.font.size = Pt(16)
                bp.font.color.rgb = DARK
                bp.space_after = Pt(14)
        elif idx == 0:
            # Title slide extra styling
            intro_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.7), Inches(3.5))
            itf = intro_box.text_frame
            itf.word_wrap = True
            p1 = itf.paragraphs[0]
            p1.text = "Production Web Application Demo & Technical Defense"
            p1.font.size = Pt(20)
            p1.font.bold = True
            p1.font.color.rgb = DARK
            p1.space_after = Pt(10)

            p2 = itf.add_paragraph()
            p2.text = "Live Deployment: https://ai-employee-tracking-system.onrender.com\nDatabase: TiDB Cloud Starter (MySQL 8.0 Protocol with TLS)\nCodebase: github.com/Vijju2005/AI-Employee-Tracking-System"
            p2.font.size = Pt(15)
            p2.font.color.rgb = GRAY

    output_path = "Employee_Break_Tracking_System_Presentation.pptx"
    prs.save(output_path)
    print(f"[+] Presentation saved successfully: {output_path}")

if __name__ == "__main__":
    create_presentation()
