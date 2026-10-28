"""
EBT System – Reports Routes
"""

import os
import io
from datetime import date, datetime
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required, current_user

from app.extensions import db
from app.models import Attendance, Break, LeaveRequest, Employee, UserRole
from app.time_utils import format_local_time, format_local_date

reports_bp = Blueprint("reports", __name__, url_prefix="/reports")


@reports_bp.route("/")
@login_required
def index():
    employees = Employee.query.order_by(Employee.full_name).all() if current_user.is_admin else []
    return render_template("reports/index.html", employees=employees)


@reports_bp.route("/generate", methods=["POST"])
@login_required
def generate():
    report_type = request.form.get("report_type", "attendance")
    fmt = request.form.get("format", "pdf")
    from_date_str = request.form.get("from_date")
    to_date_str = request.form.get("to_date")
    employee_id = request.form.get("employee_id")

    from_date = date.fromisoformat(from_date_str) if from_date_str else None
    to_date = date.fromisoformat(to_date_str) if to_date_str else None

    # Build query
    if report_type == "attendance":
        query = Attendance.query.join(Employee)
        if not current_user.is_admin and current_user.employee:
            query = query.filter(Attendance.employee_id == current_user.employee.id)
        elif employee_id and current_user.is_admin:
            query = query.filter(Attendance.employee_id == int(employee_id))
        if from_date:
            query = query.filter(Attendance.date >= from_date)
        if to_date:
            query = query.filter(Attendance.date <= to_date)
        records = query.order_by(Attendance.date.desc()).all()
        data = _attendance_to_rows(records)
        headers = ["Date", "Employee", "Check In", "Check Out", "Net Hours (h)", "Break Time (m)", "Status"]
        title = "Attendance Report"

    elif report_type == "breaks":
        query = Break.query.join(Employee, Break.employee_id == Employee.id)
        if not current_user.is_admin and current_user.employee:
            query = query.filter(Break.employee_id == current_user.employee.id)
        elif employee_id and current_user.is_admin:
            query = query.filter(Break.employee_id == int(employee_id))
        if from_date:
            query = query.filter(Break.start_time >= from_date)
        if to_date:
            query = query.filter(Break.start_time <= (to_date.isoformat() + " 23:59:59"))
        records = query.order_by(Break.start_time.desc()).all()
        data = _breaks_to_rows(records)
        headers = ["Date", "Employee", "Category", "Start", "End", "Duration (m)"]
        title = "Break Report"

    elif report_type == "leave":
        query = LeaveRequest.query.join(Employee)
        if not current_user.is_admin and current_user.employee:
            query = query.filter(LeaveRequest.employee_id == current_user.employee.id)
        elif employee_id and current_user.is_admin:
            query = query.filter(LeaveRequest.employee_id == int(employee_id))
        records = query.order_by(LeaveRequest.start_date.desc()).all()
        data = _leave_to_rows(records)
        headers = ["Employee", "Leave Type", "Start", "End", "Day Type", "Status", "Reason"]
        title = "Leave Report"

    elif report_type == "working_hours":
        query = Attendance.query.join(Employee)
        if not current_user.is_admin and current_user.employee:
            query = query.filter(Attendance.employee_id == current_user.employee.id)
        elif employee_id and current_user.is_admin:
            query = query.filter(Attendance.employee_id == int(employee_id))
        if from_date:
            query = query.filter(Attendance.date >= from_date)
        if to_date:
            query = query.filter(Attendance.date <= to_date)
        records = query.order_by(Attendance.date.desc()).all()
        data = _working_hours_to_rows(records)
        headers = ["Date", "Employee", "Dept", "Gross Hours (h)", "Break Time (m)", "Net Hours (h)", "Status"]
        title = "Working Hours Summary"

    else:
        data, headers, title = [], [], "Report"

    try:
        if fmt == "excel":
            output = _generate_excel(title, headers, data)
            filename = f"ebt_{report_type}_{date.today()}.xlsx"
            return send_file(
                io.BytesIO(output), as_attachment=True, download_name=filename,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            output = _generate_pdf(title, headers, data)
            filename = f"ebt_{report_type}_{date.today()}.pdf"
            return send_file(
                io.BytesIO(output), as_attachment=True, download_name=filename,
                mimetype="application/pdf"
            )
    except Exception as e:
        flash(f"Report generation failed: {e}", "danger")
        return redirect(url_for("reports.index"))


# ── Helpers ──────────────────────────────────────────────────────────────────

def _working_hours_to_rows(records):
    rows = []
    for r in records:
        dept = r.employee.department.name if (r.employee and r.employee.department) else "—"
        rows.append([
            str(r.date),
            r.employee.full_name if r.employee else "—",
            dept,
            f"{r.total_duration_minutes / 60:.1f}",
            f"{r.break_duration_minutes:.0f}",
            f"{r.net_working_minutes / 60:.1f}",
            r.status.value if r.status else "—",
        ])
    return rows


def _attendance_to_rows(records):
    rows = []
    for r in records:
        rows.append([
            format_local_date(r.date),
            r.employee.full_name if r.employee else "—",
            format_local_time(r.check_in, "%I:%M %p"),
            format_local_time(r.check_out, "%I:%M %p"),
            f"{r.net_working_minutes / 60:.1f}",
            f"{r.break_duration_minutes:.0f}",
            r.status.value if r.status else "—",
        ])
    return rows


def _breaks_to_rows(records):
    return [
        [
            format_local_date(r.start_time),
            r.employee.full_name if r.employee else "—",
            r.category.value if r.category else "—",
            format_local_time(r.start_time, "%I:%M %p"),
            format_local_time(r.end_time, "%I:%M %p"),
            f"{r.duration_minutes:.0f}",
        ]
        for r in records
    ]


def _leave_to_rows(records):
    return [
        [
            r.employee.full_name if r.employee else "—",
            r.leave_type.value if r.leave_type else "—",
            str(r.start_date),
            str(r.end_date),
            r.day_type.value if r.day_type else "—",
            r.status.value if r.status else "—",
            (r.reason or "")[:80],
        ]
        for r in records
    ]


def _generate_excel(title, headers, data):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    navy = "163A5F"
    # Title row
    ws.merge_cells(f"A1:{chr(64 + len(headers))}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")

    # Data rows
    for row_idx, row in enumerate(data, 3):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        if row_idx % 2 == 0:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill("solid", fgColor="EFF6FF")

    # Auto-width
    from openpyxl.utils import get_column_letter
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        max_len = 8
        for row_idx in range(2, len(data) + 3):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_pdf(title, headers, data):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=0.5*inch, rightMargin=0.5*inch,
                            topMargin=0.5*inch, bottomMargin=0.5*inch)

    styles = getSampleStyleSheet()
    navy = colors.HexColor("#163A5F")
    blue = colors.HexColor("#2563EB")

    elements = []
    elements.append(Paragraph(title, ParagraphStyle(
        "Title", parent=styles["Heading1"],
        fontSize=16, textColor=navy, spaceAfter=6
    )))
    elements.append(Paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    ))
    elements.append(Spacer(1, 12))

    table_data = [headers] + data
    col_count = len(headers)
    col_width = (landscape(A4)[0] - inch) / col_count

    t = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    doc.build(elements)
    return buf.getvalue()
